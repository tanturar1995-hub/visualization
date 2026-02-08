"""
Optional reproducibility helper:
This script regenerates the Excel workbook with the 100% stacked column chart.

Run:
  python3 code/make_excel_viz2.py
"""

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Alignment

ROOT = Path(__file__).resolve().parent.parent
PCT_CSV = ROOT / "data" / "crosstab_region_programtype_percent.csv"
OUT_XLSX = ROOT / "viz2_excel_programtype_by_region.xlsx"

def main() -> None:
    pct_df = pd.read_csv(PCT_CSV)
    wb = Workbook()
    ws = wb.active
    ws.title = "PercentTable"

    for r in dataframe_to_rows(pct_df, index=False, header=True):
        ws.append(r)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    chart.title = "Program Type Composition by Region (100% Stacked)"
    chart.y_axis.title = "Percent"
    chart.x_axis.title = "Region"

    max_row = ws.max_row
    max_col = ws.max_column
    data_ref = Reference(ws, min_col=2, min_row=1, max_col=max_col, max_row=max_row)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    ws.add_chart(chart, "H2")
    wb.save(OUT_XLSX)

if __name__ == "__main__":
    main()
